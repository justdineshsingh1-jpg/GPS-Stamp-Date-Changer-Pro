"""
GPS Stamp Date Changer Pro - Reporting & Export Module
Generates comprehensive Excel and CSV processing reports, failed item spreadsheets,
and packages ZIP archives with processed photographs.
"""

import os
import csv
import zipfile
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from importer import sanitize_url_for_logging


def generate_excel_report(job_id: str, items: List[Dict[str, Any]], out_path: str) -> str:
    """
    Creates a styled, professional Excel spreadsheet report containing
    full audit details for every processed photo.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Processing Results"
    ws.views.sheetView[0].showGridLines = True

    # Title Banner
    ws.merge_cells("A1:I1")
    title_cell = ws["A1"]
    title_cell.value = "GPS Stamp Date Changer Pro - Batch Audit Report"
    title_cell.font = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # Headers
    headers = [
        "Image ID",
        "Original URL",
        "Assigned Date",
        "Assigned Weekday",
        "Original Detected Timestamp",
        "Updated Timestamp",
        "Processing Status",
        "Output Filename",
        "Error / Notes"
    ]
    ws.append(headers)
    ws.row_dimensions[2].height = 24

    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style='thin', color="CBD5E1"),
        right=Side(style='thin', color="CBD5E1"),
        top=Side(style='thin', color="CBD5E1"),
        bottom=Side(style='thin', color="CBD5E1")
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(2, col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # Status fills
    status_fills = {
        "completed": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"), # Light green
        "failed": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),    # Light red
        "needs_review": PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid"), # Light yellow
    }

    status_fonts = {
        "completed": Font(name="Segoe UI", size=9, bold=True, color="166534"),
        "failed": Font(name="Segoe UI", size=9, bold=True, color="991B1B"),
        "needs_review": Font(name="Segoe UI", size=9, bold=True, color="854D0E"),
    }

    # Data Rows
    for row_idx, item in enumerate(items, start=3):
        ws.row_dimensions[row_idx].height = 20
        status = item.get("status", "unknown")
        
        row_vals = [
            item.get("image_id", f"IMG_{item.get('row_id', row_idx)}"),
            sanitize_url_for_logging(item.get("original_url", "")),
            item.get("assigned_date", ""),
            item.get("assigned_weekday", ""),
            item.get("original_time", ""),
            item.get("updated_timestamp", ""),
            status.upper(),
            item.get("edited_file", ""),
            item.get("error_message", "")
        ]

        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row_idx, col_idx, val)
            cell.font = Font(name="Segoe UI", size=9)
            cell.border = thin_border
            
            if col_idx in (1, 3, 4, 7):
                cell.alignment = center_align
            else:
                cell.alignment = left_align

            if col_idx == 7: # Status column styling
                cell.fill = status_fills.get(status, PatternFill())
                cell.font = status_fonts.get(status, Font(name="Segoe UI", size=9, bold=True))

    # Auto-adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col_letter].width = min(50, max(12, max_len + 3))

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    return out_path


def generate_csv_report(job_id: str, items: List[Dict[str, Any]], out_path: str) -> str:
    """Generates a standard CSV report."""
    headers = [
        "Image ID",
        "Original URL",
        "Assigned Date",
        "Assigned Weekday",
        "Original Detected Timestamp",
        "Updated Timestamp",
        "Processing Status",
        "Output Filename",
        "Error Details"
    ]
    
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    with open(out_path, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        for item in items:
            writer.writerow([
                item.get("image_id", ""),
                sanitize_url_for_logging(item.get("original_url", "")),
                item.get("assigned_date", ""),
                item.get("assigned_weekday", ""),
                item.get("original_time", ""),
                item.get("updated_timestamp", ""),
                item.get("status", ""),
                item.get("edited_file", ""),
                item.get("error_message", "")
            ])
            
    return out_path


def generate_failed_records_report(items: List[Dict[str, Any]], out_path: str) -> str:
    """
    Exports only the failed and needs_review records as a reusable spreadsheet
    so the user can re-import or re-run them directly.
    """
    failed_items = [i for i in items if i.get("status") in ("failed", "needs_review")]
    headers = ["Image ID", "Photo URL", "Failure Reason", "Suggested Date"]

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Failed Records"
    ws.append(headers)

    for item in failed_items:
        ws.append([
            item.get("image_id", ""),
            item.get("original_url", ""),
            item.get("error_message", "Unknown error"),
            item.get("assigned_date", "")
        ])

    wb.save(out_path)
    return out_path


def build_zip_package(
    job_id: str,
    job_dir: str,
    edited_dir: str,
    items: List[Dict[str, Any]],
    excel_path: Optional[str],
    csv_path: Optional[str],
    selected_row_ids: Optional[List[int]] = None
) -> str:
    """
    Packages successfully edited images and the audit reports into a compressed ZIP file.
    Supports packaging all images or only selected image row_ids.
    """
    zip_fname = "GPS_Stamp_Date_Changer_Package.zip" if not selected_row_ids else "GPS_Stamp_Selected_Photos.zip"
    zip_path = os.path.join(job_dir, zip_fname)

    with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zf:
        # Include Reports
        if excel_path and os.path.exists(excel_path):
            zf.write(excel_path, arcname="GPS_Stamp_Audit_Report.xlsx")
        if csv_path and os.path.exists(csv_path):
            zf.write(csv_path, arcname="GPS_Stamp_Audit_Report.csv")

        # Include Photos
        target_items = items
        if selected_row_ids:
            sel_set = set(selected_row_ids)
            target_items = [i for i in items if i.get("row_id") in sel_set]

        for item in target_items:
            if item.get("status") == "completed":
                f_name = item.get("edited_file")
                if f_name:
                    p = os.path.join(edited_dir, f_name)
                    if os.path.exists(p):
                        zf.write(p, arcname=f"photos/{f_name}")

    return zip_path
